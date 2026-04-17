"""
Ingestion Service — parses, chunks, embeds and stores documents.
Also registers every document in the document_registry for Layer 3 retrieval.
"""
import uuid
import logging
from pathlib import Path
from typing import List, Dict, Any

from qdrant_client.models import PointStruct
from langchain_text_splitters import RecursiveCharacterTextSplitter

from app.core.config import settings
from app.core.vector_store import get_qdrant, ensure_collection
from app.services.embedding import get_embeddings
from app.services.storage import upload_to_minio

log = logging.getLogger("intellirag.ingestion")


async def ingest_document(
    file_bytes: bytes,
    filename: str,
    tenant_slug: str,
    document_id: str,
    uploaded_by: str = "system",
    extra_metadata: dict = None,
    collection_name: str = None,
) -> Dict[str, Any]:
    """
    Full ingestion pipeline:
    1. Save to MinIO
    2. Parse text from document
    3. Split into chunks
    4. Embed each chunk
    5. Store in Qdrant with metadata
    6. Register in document_registry (Layer 3)
    """

    # 1. Store raw file in MinIO
    object_key = f"{tenant_slug}/{document_id}/{filename}"
    await upload_to_minio(file_bytes, object_key, filename)

    # 2. Parse text
    text = await parse_document(file_bytes, filename)
    if not text.strip():
        raise ValueError(f"Could not extract text from {filename}")

    # Prepend file date so embeddings encode the date context
    if extra_metadata and extra_metadata.get("file_date"):
        fd = extra_metadata["file_date"][:10]
        text = f"[Document Date: {fd}]\n\n{text}"

    # 3. Split into chunks
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.CHUNK_SIZE,
        chunk_overlap=settings.CHUNK_OVERLAP,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_text(text)

    # 4. Embed all chunks
    embeddings = await get_embeddings(chunks)

    # 5. Upsert into Qdrant
    target_collection = collection_name or await ensure_collection(tenant_slug)
    points = [
        PointStruct(
            id=str(uuid.uuid4()),
            vector=embedding,
            payload={
                "document_id": document_id,
                "tenant_slug": tenant_slug,
                "filename": filename,
                "chunk_index": i,
                "chunk_total": len(chunks),
                "text": chunk,
                "uploaded_by": uploaded_by,
                **(extra_metadata or {}),
            },
        )
        for i, (chunk, embedding) in enumerate(zip(chunks, embeddings))
    ]
    client = get_qdrant()
    await client.upsert(collection_name=target_collection, points=points)

    # 6. Register in document_registry (Layer 3)
    try:
        from app.services.registry import upsert_registry_entry
        file_ext = Path(filename).suffix.lower().lstrip(".")
        file_path = (extra_metadata or {}).get("file_path", filename)
        await upsert_registry_entry(
            document_id=document_id,
            filename=filename,
            collection=target_collection,
            file_type=file_ext,
            file_path=file_path,
            chunk_count=len(chunks),
            text_sample=text[:3000],
        )
    except Exception as e:
        # Registry failure should never block ingestion
        log.warning(f"[Ingestion] Registry update failed for {filename}: {e}")

    return {
        "document_id": document_id,
        "filename": filename,
        "chunks_created": len(chunks),
        "collection": target_collection,
    }


async def parse_document(file_bytes: bytes, filename: str) -> str:
    """Parse text from various file types."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        return await _parse_pdf(file_bytes)
    elif ext in (".docx", ".doc"):
        return await _parse_docx(file_bytes)
    elif ext in (".xlsx", ".xls"):
        return await _parse_xlsx(file_bytes)
    elif ext in (".pptx", ".ppt"):
        return await _parse_pptx(file_bytes)
    elif ext in (".txt", ".md", ".csv"):
        return file_bytes.decode("utf-8", errors="ignore")
    else:
        return file_bytes.decode("utf-8", errors="ignore")


async def _parse_pdf(file_bytes: bytes) -> str:
    import io
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            pages.append(t)
    return "\n\n".join(pages)


async def _parse_docx(file_bytes: bytes) -> str:
    import io
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


async def _parse_xlsx(file_bytes: bytes) -> str:
    """Parse Excel — each sheet becomes a named section."""
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    sections = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            values = [str(v).strip() if v is not None else "" for v in row]
            if any(v for v in values):
                rows.append("  |  ".join(values))
        if rows:
            sections.append(f"Sheet: {sheet_name}\n" + "\n".join(rows))
    return "\n\n".join(sections)


async def _parse_pptx(file_bytes: bytes) -> str:
    """Parse PowerPoint — each slide becomes a named section."""
    import io
    from pptx import Presentation
    prs = Presentation(io.BytesIO(file_bytes))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                texts.append(shape.text.strip())
        if texts:
            slides.append(f"Slide {i}:\n" + "\n".join(texts))
    return "\n\n".join(slides)


async def delete_document_chunks(document_id: str, tenant_slug: str):
    """Remove all Qdrant points and registry entry for a document."""
    from qdrant_client.models import Filter, FieldCondition, MatchValue
    client = get_qdrant()
    collection_name = f"intellirag_{tenant_slug}"
    await client.delete(
        collection_name=collection_name,
        points_selector=Filter(
            must=[FieldCondition(key="document_id", match=MatchValue(value=document_id))]
        ),
    )
    # Also remove from registry
    try:
        from app.services.registry import remove_registry_entry
        await remove_registry_entry(document_id, collection_name)
    except Exception as e:
        log.warning(f"[Ingestion] Registry removal failed: {e}")
